"""
Módulo de Geração de Relatórios Executivos em DOCX e Planilhamento de Ensaios
"""

import csv
import logging
import os
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from config import LOGS_DIR, REPORTS_DIR, TRIALS_LOG_FILE

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)


def _set_cell_background(cell, hex_color: str):
    """Aplica cor de fundo a uma célula de tabela docx."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Ajusta espaçamento interno da célula."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


class ReportGenerator:
    """
    Gerador de relatórios executivos em formato Word (.docx) e
    registrador automático de ensaios analíticos em planilha CSV.
    """

    @staticmethod
    def log_trial(trial_data: Dict[str, Any]) -> str:
        """
        Registra o ensaio executado no histórico em formato CSV.
        """
        file_exists = TRIALS_LOG_FILE.exists()
        fieldnames = [
            "Data_Hora",
            "Ativo",
            "Ticker",
            "Moeda",
            "Periodo",
            "Populacao",
            "Geracoes",
            "Taxa_Mutacao",
            "Taxa_Crossover",
            "Lookback",
            "Horizonte_Dias",
            "Ultimo_Preco_Real",
            "Preco_Projetado_Final",
            "Variacao_Esperada_Pct",
            "Tendencia",
            "RMSE",
            "MAE",
            "MAPE_Pct",
            "R2",
            "Acuracia_Direcional_Pct"
        ]

        with open(TRIALS_LOG_FILE, mode="a", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            if not file_exists:
                writer.writeheader()
            writer.writerow(trial_data)

        logger.info(f"Ensaio registrado com sucesso em: {TRIALS_LOG_FILE}")
        return str(TRIALS_LOG_FILE)

    @staticmethod
    def generate_docx_report(
        asset_name: str,
        ticker: str,
        currency: str,
        period_str: str,
        ga_config: Dict[str, Any],
        prediction_results: Dict[str, Any],
        ga_history: Dict[str, Any],
        chart_paths: Dict[str, str],
        output_filepath: Optional[str] = None
    ) -> str:
        """
        Gera um relatório executivo profissional em formato .docx contendo
        tabelas formatadas, sumário estatístico e gráficos embutidos.
        """
        doc = Document()

        # Configuração de Margens (1 polegada)
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Paleta de Cores
        primary_color = RGBColor(16, 44, 87)       # Azul escuro executivo
        secondary_color = RGBColor(53, 95, 142)    # Azul médio
        accent_color = RGBColor(46, 125, 50)       # Verde sucesso
        header_hex = "102C57"
        row_alt_hex = "F4F6F9"
        white_hex = "FFFFFF"

        # 1. Cabeçalho e Título do Documento
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_cat = title_p.add_run("RELATÓRIO EXECUTIVO DE INTELIGÊNCIA ARTIFICIAL\n")
        run_cat.font.size = Pt(10)
        run_cat.font.bold = True
        run_cat.font.color.rgb = secondary_color

        run_title = title_p.add_run("Previsão e Otimização de Ativos Financeiros via Algoritmo Genético")
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = primary_color

        meta_p = doc.add_paragraph()
        meta_p.paragraph_format.space_after = Pt(16)
        now_str = datetime.now().strftime("%d/%m/%Y às %H:%M:%S")
        run_meta = meta_p.add_run(f"Ativo Analisado: {asset_name} ({ticker}) | Data de Emissão: {now_str}")
        run_meta.font.size = Pt(10)
        run_meta.font.italic = True
        run_meta.font.color.rgb = RGBColor(100, 100, 100)

        # Divisor visual
        doc.add_paragraph("―" * 55).paragraph_format.space_after = Pt(14)

        # 2. Resumo Executivo
        metrics = prediction_results["metrics"]
        h1 = doc.add_heading("1. Resumo Executivo e Projeção Futura", level=1)
        h1.style.font.color.rgb = primary_color

        summary_p = doc.add_paragraph()
        summary_p.paragraph_format.line_spacing = 1.25
        summary_p.paragraph_format.space_after = Pt(14)

        last_p = metrics['ultimo_preco_real']
        proj_p = metrics['preco_projetado_final']
        var_p = metrics['variacao_esperada_pct']
        trend_str = metrics['tendencia_esperada']
        horizon_days = metrics.get('horizonte_dias', 5)

        summary_text = (
            f"O presente relatório consolida o resultado do treinamento e otimização por Algoritmo Genético "
            f"para o ativo {asset_name} ({ticker}), cotado em {currency}. Com base na série temporal do período "
            f"({period_str}) e na evolução dos pesos preditivos, a projeção para os próximos {horizon_days} períodos úteis "
            f"aponta para um preço estimado de {currency} {proj_p:,.4f}, comparado ao último fechamento de {currency} {last_p:,.4f}. "
            f"Isso representa uma variação esperada de {var_p:+.2f}%, com viés de tendência classificado como: {trend_str}."
        )
        summary_p.add_run(summary_text)

        # 3. Tabela de Métricas de Validação e Desempenho
        h2 = doc.add_heading("2. Métricas Estatísticas de Aderência do Modelo", level=1)
        h2.style.font.color.rgb = primary_color

        table_metrics = doc.add_table(rows=6, cols=3)
        table_metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_metrics.autofit = False

        headers = ["Métrica Estatística", "Valor Obtido", "Interpretação Técnica"]
        hdr_cells = table_metrics.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            _set_cell_background(hdr_cells[i], header_hex)
            _set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

        metric_rows = [
            ("RMSE (Raiz do Erro Quadrático Médio)", f"{metrics['rmse']:,.4f}", "Erro médio ponderado em magnitude absoluta"),
            ("MAE (Erro Absoluto Médio)", f"{metrics['mae']:,.4f}", "Desvio médio linear entre o real e o previsto"),
            ("MAPE (Erro Percentual Médio)", f"{metrics['mape']:,.2f}%", "Erro relativo médio em termos percentuais"),
            ("Coeficiente de Determinação (R²)", f"{metrics['r2']:,.4f}", "Capacidade explicativa da variabilidade da série"),
            ("Acurácia Direcional de Sinal", f"{metrics['acuracia_direcional']:,.2f}%", "Taxa de acerto na previsão do sentido do movimento (Alta/Baixa)")
        ]

        for row_idx, (m_name, m_val, m_desc) in enumerate(metric_rows, start=1):
            row_cells = table_metrics.rows[row_idx].cells
            row_cells[0].text = m_name
            row_cells[1].text = m_val
            row_cells[2].text = m_desc

            bg_color = row_alt_hex if row_idx % 2 == 0 else white_hex
            for col_idx, cell in enumerate(row_cells):
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                p = cell.paragraphs[0]
                if col_idx == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if len(p.runs) > 0:
                        p.runs[0].font.bold = True
                for run in p.runs:
                    run.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 4. Tabela de Parâmetros do Algoritmo Genético
        h3 = doc.add_heading("3. Hiperparâmetros do Algoritmo Genético Empregado", level=1)
        h3.style.font.color.rgb = primary_color

        table_ga = doc.add_table(rows=7, cols=2)
        table_ga.alignment = WD_TABLE_ALIGNMENT.CENTER
        ga_hdr = table_ga.rows[0].cells
        ga_hdr[0].text = "Parâmetro Evolutivo"
        ga_hdr[1].text = "Valor Configurado"
        for cell in ga_hdr:
            _set_cell_background(cell, header_hex)
            _set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

        ga_rows = [
            ("Tamanho da População", str(ga_config.get("population_size", 60))),
            ("Número de Gerações", str(ga_config.get("generations", 35))),
            ("Taxa de Recombinação (Crossover)", f"{ga_config.get('crossover_rate', 0.85) * 100:.1f}%"),
            ("Taxa de Mutação Gênica Adaptativa", f"{ga_config.get('mutation_rate', 0.15) * 100:.1f}%"),
            ("Proporção de Elitismo", f"{ga_config.get('elitism_ratio', 0.08) * 100:.1f}%"),
            ("Janela de Lookback Temporal (Lags)", f"{ga_config.get('lookback_window', 10)} períodos")
        ]

        for row_idx, (p_name, p_val) in enumerate(ga_rows, start=1):
            row_cells = table_ga.rows[row_idx].cells
            row_cells[0].text = p_name
            row_cells[1].text = p_val

            bg_color = row_alt_hex if row_idx % 2 == 0 else white_hex
            for col_idx, cell in enumerate(row_cells):
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=90, bottom=90, left=150, right=150)
                p = cell.paragraphs[0]
                if col_idx == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if len(p.runs) > 0:
                        p.runs[0].font.bold = True
                for run in p.runs:
                    run.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 5. Tabela de Projeção Detalhada Passo a Passo
        h4 = doc.add_heading(f"4. Projeção Passo a Passo para os Próximos {horizon_days} Dias", level=1)
        h4.style.font.color.rgb = primary_color

        df_forecast = prediction_results["forecast_df"]
        table_fore = doc.add_table(rows=len(df_forecast) + 1, cols=4)
        table_fore.alignment = WD_TABLE_ALIGNMENT.CENTER

        fore_headers = ["Data Futura", f"Preço Previsto ({currency})", "Limite Inferior (95%)", "Limite Superior (95%)"]
        for col_idx, text in enumerate(fore_headers):
            cell = table_fore.rows[0].cells[col_idx]
            cell.text = text
            _set_cell_background(cell, header_hex)
            _set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

        for row_idx, (date_val, row) in enumerate(df_forecast.iterrows(), start=1):
            cells = table_fore.rows[row_idx].cells
            date_formatted = date_val.strftime("%d/%m/%Y")
            cells[0].text = date_formatted
            cells[1].text = f"{row['Preco_Projetado']:,.4f}"
            cells[2].text = f"{row['Limite_Inferior']:,.4f}"
            cells[3].text = f"{row['Limite_Superior']:,.4f}"

            bg_color = row_alt_hex if row_idx % 2 == 0 else white_hex
            for col_idx, cell in enumerate(cells):
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=90, bottom=90, left=150, right=150)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if col_idx == 1 and len(p.runs) > 0:
                    p.runs[0].font.bold = True
                for run in p.runs:
                    run.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(16)

        # 6. Gráficos Embutidos em Alta Resolução
        h5 = doc.add_heading("5. Visualização Gráfica e Curva de Aprendizado", level=1)
        h5.style.font.color.rgb = primary_color

        if "prediction_chart" in chart_paths and os.path.exists(chart_paths["prediction_chart"]):
            p_img1_title = doc.add_paragraph()
            r1 = p_img1_title.add_run("Figura 1: Série Histórica, Ajuste da Inteligência Artificial e Projeção Futura")
            r1.font.bold = True
            r1.font.size = Pt(10.5)
            doc.add_picture(chart_paths["prediction_chart"], width=Inches(6.2))
            doc.add_paragraph().paragraph_format.space_after = Pt(12)

        if "convergence_chart" in chart_paths and os.path.exists(chart_paths["convergence_chart"]):
            p_img2_title = doc.add_paragraph()
            r2 = p_img2_title.add_run("Figura 2: Curva de Convergência Evolutiva do Algoritmo Genético (Fitness por Geração)")
            r2.font.bold = True
            r2.font.size = Pt(10.5)
            doc.add_picture(chart_paths["convergence_chart"], width=Inches(6.2))
            doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 7. Considerações e Recomendações Técnicas
        h6 = doc.add_heading("6. Conclusões e Metodologia", level=1)
        h6.style.font.color.rgb = primary_color
        p_conc = doc.add_paragraph()
        p_conc.paragraph_format.line_spacing = 1.25
        p_conc.add_run(
            "O Algoritmo Genético atingiu convergência estável no espaço de busca, combinando mecanismos "
            "de seleção por torneio e mutação adaptativa decrescente para mitigar o aprisionamento em mínimos locais. "
            "A incorporação de múltiplos indicadores técnicos (RSI, Bandas de Bollinger, Médias Móveis e MACD) "
            "em conjunto com defasagens autorregressivas (lags) confere resiliência e estabilidade preditiva às "
            "séries temporais financeiras e de commodities. Recomenda-se a re-otimização periódica (ensaios contínuos) "
            "conforme novas cotações sejam consolidadas no mercado."
        )

        # Salvar documento
        if not output_filepath:
            from data_fetcher import DataFetcher
            nome_base = DataFetcher.obter_caminho_arquivo_excel(ticker).stem

            if prediction_results and "history_df" in prediction_results and not prediction_results["history_df"].empty:
                idx = prediction_results["history_df"].index
                dt_ini = pd.to_datetime(idx.min()).strftime("%d-%m-%Y")
                dt_fim = pd.to_datetime(idx.max()).strftime("%d-%m-%Y")
            else:
                agora = datetime.now()
                dt_fim = agora.strftime("%d-%m-%Y")
                dt_ini = (agora - timedelta(days=365)).strftime("%d-%m-%Y")

            output_filepath = str(REPORTS_DIR / f"{nome_base} {dt_ini} a {dt_fim}.docx")

        doc.save(output_filepath)
        logger.info(f"Relatório executivo DOCX salvo com sucesso em: {output_filepath}")
        return output_filepath

    @staticmethod
    def export_charts_for_report(
        prediction_results: Dict[str, Any],
        ga_history: Dict[str, Any],
        asset_name: str,
        currency: str,
        temp_dir: Optional[Path] = None
    ) -> Dict[str, str]:
        """
        Exporta gráficos em formato PNG de alta resolução para inclusão no DOCX.
        """
        target_dir = temp_dir or REPORTS_DIR
        target_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")

        pred_chart_path = str(target_dir / f"chart_pred_{ts}.png")
        conv_chart_path = str(target_dir / f"chart_conv_{ts}.png")

        # 1. Gráfico de Predição e Projeção
        plt.style.use("seaborn-v0_8-whitegrid" if "seaborn-v0_8-whitegrid" in plt.style.available else "default")
        fig1, ax1 = plt.subplots(figsize=(10, 5), dpi=300)

        df_history = prediction_results["history_df"]
        df_forecast = prediction_results["forecast_df"]

        # Plota os últimos 90 dias de histórico para clareza visual
        plot_history = df_history.tail(90)

        ax1.plot(plot_history.index, plot_history["Preco_Real"], label="Preço Real (Mercado)", color="#1F77B4", lw=2.0)
        ax1.plot(plot_history.index, plot_history["Preco_Previsto_IA"], label="Ajuste Algoritmo Genético", color="#FF7F0E", lw=1.8, linestyle="--")

        # Linha e área de projeção futura
        last_hist_date = plot_history.index[-1]
        last_hist_price = plot_history["Preco_Real"].iloc[-1]

        proj_dates = [last_hist_date] + list(df_forecast.index)
        proj_prices = [last_hist_price] + list(df_forecast["Preco_Projetado"])
        proj_lower = [last_hist_price] + list(df_forecast["Limite_Inferior"])
        proj_upper = [last_hist_price] + list(df_forecast["Limite_Superior"])

        ax1.plot(proj_dates, proj_prices, label="Projeção Futura (IA)", color="#2CA02C", lw=2.2, marker="o", markersize=4)
        ax1.fill_between(proj_dates, proj_lower, proj_upper, color="#2CA02C", alpha=0.18, label="Intervalo de Confiança (95%)")

        ax1.set_title(f"Série Temporal e Projeção Futura - {asset_name}", fontsize=13, fontweight="bold", pad=12)
        ax1.set_xlabel("Data", fontsize=10)
        ax1.set_ylabel(f"Preço ({currency})", fontsize=10)
        ax1.legend(loc="best", frameon=True)
        fig1.tight_layout()
        fig1.savefig(pred_chart_path, dpi=300)
        plt.close(fig1)

        # 2. Gráfico de Convergência do Algoritmo Genético
        fig2, ax2 = plt.subplots(figsize=(10, 4.5), dpi=300)
        generations = ga_history["generation"]
        best_fitness = ga_history["best_fitness"]
        avg_fitness = ga_history["avg_fitness"]

        ax2.plot(generations, best_fitness, label="Melhor Fitness da Geração", color="#D62728", lw=2.2)
        ax2.plot(generations, avg_fitness, label="Fitness Médio da População", color="#1F77B4", lw=1.8, linestyle=":")

        ax2.set_title("Curva de Aprendizado e Convergência Evolutiva (GA)", fontsize=13, fontweight="bold", pad=12)
        ax2.set_xlabel("Geração", fontsize=10)
        ax2.set_ylabel("Valor de Aptidão (Fitness)", fontsize=10)
        ax2.legend(loc="best", frameon=True)
        fig2.tight_layout()
        fig2.savefig(conv_chart_path, dpi=300)
        plt.close(fig2)

        return {
            "prediction_chart": pred_chart_path,
            "convergence_chart": conv_chart_path
        }
